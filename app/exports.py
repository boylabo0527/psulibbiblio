"""Export the master recommendation table to xlsx, csv, pdf, docx."""
from __future__ import annotations

import io
from typing import Sequence

import pandas as pd

COLUMNS = [
    "Campus", "College", "Program", "Course",
    "Book Title", "Author", "Publication Year", "Publisher", "Number of Copies",
]


def _frame(rows: Sequence[dict]) -> pd.DataFrame:
    df = pd.DataFrame(rows, columns=COLUMNS)
    return df.sort_values(by=COLUMNS[:5], kind="stable").reset_index(drop=True)


def to_xlsx(rows: Sequence[dict]) -> bytes:
    df = _frame(rows)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Recommendations")
    return buf.getvalue()


def to_csv(rows: Sequence[dict]) -> bytes:
    return _frame(rows).to_csv(index=False).encode("utf-8")


def to_pdf(rows: Sequence[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    df = _frame(rows)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title="Library Acquisition Recommendations")
    styles = getSampleStyleSheet()
    elems = [Paragraph("Library Acquisition Recommendations", styles["Title"]), Spacer(1, 12)]
    data = [list(df.columns)] + df.astype(str).values.tolist()
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    elems.append(table)
    doc.build(elems)
    return buf.getvalue()


def to_docx(rows: Sequence[dict]) -> bytes:
    import docx
    df = _frame(rows)
    d = docx.Document()
    d.add_heading("Library Acquisition Recommendations", level=1)
    table = d.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr[i].text = col
    for _, r in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            cells[i].text = str(r[col])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
